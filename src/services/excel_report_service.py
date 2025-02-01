import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
from openpyxl.styles import PatternFill, Font, Alignment
from ..interfaces.report_interfaces import IReportGenerator
from ..core.cache_manager import CacheManager
from datetime import datetime, date
from pathlib import Path
from typing import Dict, List, Any
from openpyxl.utils import get_column_letter
from ..interfaces.shift_interfaces import Schedule, Shift, ShiftType
from ..interfaces.staff_interfaces import Doctor
from ..services.data_loader import StaffDataLoader

class ExcelReportService(IReportGenerator):
    """Service for generating Excel reports from schedules"""
    
    def __init__(self, staff_loader: StaffDataLoader):
        self.staff_loader = staff_loader
        self._staff_dict = self._create_staff_dict()
        self.style_header = Font(bold=True, size=12)
        self.highlight_fill = PatternFill(
            start_color='FFE6E6',
            end_color='FFE6E6',
            fill_type='solid'
        )
        self.workbook = None
        self.cache = CacheManager()
        
    def _create_staff_dict(self) -> Dict[str, Doctor]:
        """Create a lookup dictionary for staff members"""
        staff = self.staff_loader.load_all_staff()
        return {doctor.id: doctor for doctor in staff}
    
    def generate_schedule_report(self, schedule: Schedule, output_path: str) -> None:
        """Generate an Excel report for a schedule
        
        Args:
            schedule: The schedule to generate a report for
            output_path: Path to save the Excel file
        """
        # Create Excel writer
        writer = pd.ExcelWriter(output_path, engine='openpyxl')
        
        # Generate daily schedule sheet
        self._generate_daily_schedule_sheet(schedule, writer)
        
        # Generate staff summary sheet
        self._generate_staff_summary_sheet(schedule, writer)
        
        # Generate team balance sheet
        self._generate_team_balance_sheet(schedule, writer)
        
        # Generate shift distribution sheet
        self._generate_shift_distribution_sheet(schedule, writer)
        
        # Save the Excel file
        writer.close()
    
    def _generate_daily_schedule_sheet(self, schedule: Schedule, writer: pd.ExcelWriter) -> None:
        """Generate daily schedule sheet"""
        # Get date range
        all_shifts = schedule.get_all_shifts()
        if not all_shifts:
            return
            
        start_date = min(s.date.date() for s in all_shifts)
        end_date = max(s.date.date() for s in all_shifts)
        
        # Create data for each day
        data = []
        current_date = start_date
        while current_date <= end_date:
            day_shifts = schedule.get_shifts_by_date(current_date)
            for shift in day_shifts:
                doctor = self._staff_dict.get(shift.assigned_doctor_id) if shift.assigned_doctor_id else None
                data.append({
                    'Date': current_date,
                    'Day': current_date.strftime('%A'),
                    'Shift Type': shift.shift_definition.shift_type.value,
                    'Time': f"{shift.shift_definition.time_slot.start_time.strftime('%H:%M')}-{shift.shift_definition.time_slot.end_time.strftime('%H:%M')}",
                    'Doctor': doctor.name if doctor else 'UNASSIGNED',
                    'Team': f"Team {doctor.team}" if doctor else 'N/A',
                    'Gender': doctor.gender if doctor else 'N/A'
                })
            current_date = date.fromordinal(current_date.toordinal() + 1)
        
        # Create DataFrame and write to Excel
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Daily Schedule', index=False)
        
        # Format the sheet
        sheet = writer.sheets['Daily Schedule']
        for column in sheet.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = adjusted_width
    
    def _generate_staff_summary_sheet(self, schedule: Schedule, writer: pd.ExcelWriter) -> None:
        """Generate staff summary sheet"""
        data = []
        for doctor_id, doctor in self._staff_dict.items():
            shifts = schedule.get_shifts_by_doctor(doctor_id)
            total_hours = sum(
                (s.shift_definition.time_slot.end_time.hour - 
                 s.shift_definition.time_slot.start_time.hour +
                 (s.shift_definition.time_slot.end_time.minute - 
                  s.shift_definition.time_slot.start_time.minute) / 60)
                for s in shifts
            )
            late_shifts = len([
                s for s in shifts
                if s.shift_definition.time_slot.start_time.hour >= 17
            ])
            
            data.append({
                'Doctor': doctor.name,
                'Team': f"Team {doctor.team}",
                'Gender': doctor.gender,
                'Total Shifts': len(shifts),
                'Total Hours': round(total_hours, 1),
                'Late Shifts': late_shifts
            })
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Staff Summary', index=False)
    
    def _generate_team_balance_sheet(self, schedule: Schedule, writer: pd.ExcelWriter) -> None:
        """Generate team balance analysis sheet"""
        team_stats = {1: {}, 2: {}}
        shift_types = list(ShiftType)
        
        # Initialize stats
        for team in team_stats:
            for shift_type in shift_types:
                team_stats[team][shift_type.value] = 0
        
        # Count shifts by team and type
        for shift in schedule.get_all_shifts():
            if shift.assigned_doctor_id:
                doctor = self._staff_dict[shift.assigned_doctor_id]
                team_stats[doctor.team][shift.shift_definition.shift_type.value] += 1
        
        # Create DataFrame
        data = []
        for team in [1, 2]:
            row = {'Team': f"Team {team}"}
            for shift_type in shift_types:
                row[shift_type.value] = team_stats[team][shift_type.value]
            data.append(row)
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Team Balance', index=False)
    
    def _generate_shift_distribution_sheet(self, schedule: Schedule, writer: pd.ExcelWriter) -> None:
        """Generate shift distribution analysis sheet"""
        shift_types = list(ShiftType)
        data = []
        
        for doctor_id, doctor in self._staff_dict.items():
            row = {
                'Doctor': doctor.name,
                'Team': f"Team {doctor.team}",
                'Gender': doctor.gender
            }
            
            shifts = schedule.get_shifts_by_doctor(doctor_id)
            for shift_type in shift_types:
                row[shift_type.value] = len([
                    s for s in shifts
                    if s.shift_definition.shift_type == shift_type
                ])
            
            data.append(row)
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Shift Distribution', index=False)
    
    def generate_detailed_report(self, schedule: Schedule) -> Dict[str, Any]:
        """Generate a detailed analysis report of the schedule
        
        Returns:
            Dictionary containing various schedule metrics
        """
        return {
            'shift_distribution': self._analyze_shift_distribution(schedule),
            'team_balance': self._analyze_team_balance(schedule),
            'gender_compliance': self._analyze_gender_compliance(schedule),
            'preference_satisfaction': self._analyze_preference_satisfaction(schedule)
        }
    
    def _get_cached_data(self, key, generator_func, ttl=3600):
        """Get data from cache or generate if not available"""
        data = self.cache.get(key)
        if data is None:
            data = generator_func()
            self.cache.set(key, data, ttl)
        return data
    
    def _analyze_shift_distribution(self, schedule: Schedule) -> Dict[str, Any]:
        """Analyze shift distribution across staff"""
        staff_shifts = {}
        total_shifts = 0
        
        for doctor_id, doctor in self._staff_dict.items():
            shifts = schedule.get_shifts_by_doctor(doctor_id)
            staff_shifts[doctor_id] = {
                'name': doctor.name,
                'team': doctor.team,
                'total': len(shifts),
                'by_type': {}
            }
            
            # Count shifts by type
            for shift_type in ShiftType:
                type_shifts = len([
                    s for s in shifts
                    if s.shift_definition.shift_type == shift_type
                ])
                staff_shifts[doctor_id]['by_type'][shift_type.value] = type_shifts
            
            total_shifts += len(shifts)
        
        # Calculate distribution metrics
        num_staff = len(self._staff_dict)
        avg_shifts = total_shifts / num_staff if num_staff > 0 else 0
        shift_counts = [stats['total'] for stats in staff_shifts.values()]
        max_shifts = max(shift_counts) if shift_counts else 0
        min_shifts = min(shift_counts) if shift_counts else 0
        
        # Calculate equity score (0 = perfectly balanced, higher = more imbalanced)
        equity_score = (max_shifts - min_shifts) / avg_shifts if avg_shifts > 0 else 0
        
        return {
            'staff_shifts': staff_shifts,
            'total_shifts': total_shifts,
            'average_shifts': avg_shifts,
            'max_shifts': max_shifts,
            'min_shifts': min_shifts,
            'equity_score': equity_score
        }
    
    def _analyze_team_balance(self, schedule: Schedule) -> Dict[str, Any]:
        """Analyze shift balance between teams"""
        team_shifts = {1: {'total': 0}, 2: {'total': 0}}
        total_shifts = 0
        
        for shift in schedule.get_all_shifts():
            if shift.assigned_doctor_id:
                doctor = self._staff_dict[shift.assigned_doctor_id]
                team_shifts[doctor.team]['total'] += 1
                total_shifts += 1
        
        # Calculate balance metrics
        expected_per_team = total_shifts / 2
        max_deviation = max(
            abs(team_shifts[1]['total'] - expected_per_team),
            abs(team_shifts[2]['total'] - expected_per_team)
        )
        
        # Calculate balance score (0 = perfectly balanced, higher = more imbalanced)
        balance_score = max_deviation / expected_per_team if expected_per_team > 0 else 0
        
        return {
            'team_shifts': team_shifts,
            'total_shifts': total_shifts,
            'expected_per_team': expected_per_team,
            'max_deviation': max_deviation,
            'balance_score': balance_score
        }
    
    def _analyze_gender_compliance(self, schedule: Schedule) -> Dict[str, Any]:
        """Analyze gender requirement compliance"""
        total_shifts = 0
        compliant_shifts = 0
        violations = []
        
        for shift in schedule.get_all_shifts():
            total_shifts += 1
            if shift.assigned_doctor_id:
                doctor = self._staff_dict[shift.assigned_doctor_id]
                requirement = shift.shift_definition.gender_requirement
                
                if requirement == 'Male Only' and doctor.gender != 'M':
                    violations.append({
                        'shift_id': shift.id,
                        'requirement': requirement,
                        'assigned': doctor.gender
                    })
                elif requirement == 'Female Only' and doctor.gender != 'F':
                    violations.append({
                        'shift_id': shift.id,
                        'requirement': requirement,
                        'assigned': doctor.gender
                    })
                else:
                    compliant_shifts += 1
        
        return {
            'total_shifts': total_shifts,
            'compliant_shifts': compliant_shifts,
            'compliance_rate': compliant_shifts / total_shifts if total_shifts > 0 else 1.0,
            'violations': violations
        }
    
    def _analyze_preference_satisfaction(self, schedule: Schedule) -> Dict[str, Any]:
        """Analyze how well staff preferences are satisfied"""
        # TODO: Implement preference satisfaction analysis once preference system is added
        return {
            'total_preferences': 0,
            'satisfied_preferences': 0,
            'satisfaction_rate': 1.0,
            'violations': []
        }
    
    def save_report(self, filename):
        """Save the generated report to an Excel file"""
        if self.workbook:
            self.workbook.save(filename) 