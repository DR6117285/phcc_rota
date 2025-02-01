from datetime import datetime, date
import json
import openpyxl
from openpyxl.styles import Font, PatternFill
from pathlib import Path
from typing import Dict, List, Any, Optional
import pandas as pd
import logging
from openpyxl import Workbook
from src.interfaces.audit_logger import IAuditLogger, ShiftAllocation
from ..interfaces.audit_interfaces import (
    AuditEvent, AuditEventType, AuditEventSeverity
)

logger = logging.getLogger(__name__)

class AuditLogger(IAuditLogger):
    """Implementation of the audit logging system with Excel report generation"""
    
    def __init__(self, log_dir: str = "logs/allocations"):
        self.log_dir = Path(log_dir)
        self.log_dir.mkdir(parents=True, exist_ok=True)
        self._events: List[AuditEvent] = []
        self._current_timestamp = None  # For testing purposes
    
    def _create_event(self,
                     event_type: AuditEventType,
                     severity: AuditEventSeverity,
                     description: str,
                     details: Dict[str, Any]) -> AuditEvent:
        """Helper method to create an audit event"""
        timestamp = self._current_timestamp or datetime.now()
        event = AuditEvent(
            event_type=event_type,
            severity=severity,
            timestamp=timestamp,
            description=description,
            details=details
        )
        self._events.append(event)
        return event
    
    def _override_timestamp(self, timestamp: datetime) -> None:
        """Override current timestamp for testing"""
        self._current_timestamp = timestamp
    
    def log_schedule_generation(self,
                              start_date: datetime,
                              end_date: datetime,
                              num_shifts: int,
                              success: bool) -> AuditEvent:
        """Log a schedule generation event"""
        description = f"Generated schedule from {start_date.date()} to {end_date.date()}"
        if not success:
            description = f"Failed to generate schedule from {start_date.date()} to {end_date.date()}"
        
        return self._create_event(
            event_type=AuditEventType.SCHEDULE_GENERATION,
            severity=AuditEventSeverity.INFO if success else AuditEventSeverity.ERROR,
            description=description,
            details={
                "start_date": start_date.isoformat(),
                "end_date": end_date.isoformat(),
                "num_shifts": num_shifts,
                "success": success
            }
        )
    
    def log_optimization_step(self,
                            iteration: int,
                            strategy: str,
                            improvement: float,
                            metrics: Dict[str, float]) -> AuditEvent:
        """Log an optimization step"""
        description = f"Optimization step {iteration} using {strategy} strategy"
        if improvement > 0:
            description += f" (improved by {improvement:.2%})"
        
        return self._create_event(
            event_type=AuditEventType.OPTIMIZATION,
            severity=AuditEventSeverity.INFO,
            description=description,
            details={
                "iteration": iteration,
                "strategy": strategy,
                "improvement": improvement,
                "metrics": metrics
            }
        )
    
    def log_constraint_violation(self,
                               constraint_type: str,
                               shift_id: str,
                               severity: AuditEventSeverity,
                               details: Dict[str, Any]) -> AuditEvent:
        """Log a constraint violation"""
        description = f"Constraint violation ({constraint_type}) for shift {shift_id}"
        
        return self._create_event(
            event_type=AuditEventType.CONSTRAINT_VIOLATION,
            severity=severity,
            description=description,
            details={
                "constraint_type": constraint_type,
                "shift_id": shift_id,
                **details
            }
        )
    
    def log_conflict_resolution(self,
                              conflict_id: str,
                              resolution_type: str,
                              success: bool,
                              affected_shifts: List[str]) -> AuditEvent:
        """Log a conflict resolution event"""
        description = f"Conflict resolution {conflict_id} using {resolution_type}"
        if not success:
            description = f"Failed conflict resolution {conflict_id} using {resolution_type}"
        
        return self._create_event(
            event_type=AuditEventType.CONFLICT_RESOLUTION,
            severity=AuditEventSeverity.INFO if success else AuditEventSeverity.WARNING,
            description=description,
            details={
                "conflict_id": conflict_id,
                "resolution_type": resolution_type,
                "success": success,
                "affected_shifts": affected_shifts
            }
        )
    
    def log_performance_metric(self,
                             metric_name: str,
                             value: float,
                             unit: str) -> AuditEvent:
        """Log a performance metric"""
        description = f"Performance metric: {metric_name} = {value} {unit}"
        
        return self._create_event(
            event_type=AuditEventType.PERFORMANCE,
            severity=AuditEventSeverity.INFO,
            description=description,
            details={
                "metric_name": metric_name,
                "value": value,
                "unit": unit
            }
        )
    
    def get_events_by_type(self, event_type: AuditEventType) -> List[AuditEvent]:
        """Get all events of a specific type"""
        return [event for event in self._events if event.event_type == event_type]
    
    def get_events_by_severity(self, severity: AuditEventSeverity) -> List[AuditEvent]:
        """Get all events of a specific severity"""
        return [event for event in self._events if event.severity == severity]
    
    def get_events_in_timerange(self,
                              start_time: datetime,
                              end_time: datetime) -> List[AuditEvent]:
        """Get all events within a specific time range"""
        return [
            event for event in self._events
            if start_time <= event.timestamp <= end_time
        ]
    
    def export_to_excel(self, filepath: str) -> None:
        """Export audit trail to Excel file"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Audit Trail"
            
            # Define headers
            headers = ["Timestamp", "Type", "Severity", "Description", "Details"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col)
                cell.value = header
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            # Add events
            for row, event in enumerate(self._events, 2):
                ws.cell(row=row, column=1).value = event.timestamp.isoformat()
                ws.cell(row=row, column=2).value = event.event_type.name
                ws.cell(row=row, column=3).value = event.severity.name
                ws.cell(row=row, column=4).value = event.description
                ws.cell(row=row, column=5).value = str(event.details)
                
                # Color-code severity
                severity_cell = ws.cell(row=row, column=3)
                if event.severity == AuditEventSeverity.ERROR:
                    severity_cell.fill = PatternFill(start_color="FFB6B6", end_color="FFB6B6", fill_type="solid")
                elif event.severity == AuditEventSeverity.WARNING:
                    severity_cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
                elif event.severity == AuditEventSeverity.CRITICAL:
                    severity_cell.fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
            
            # Adjust column widths
            for column in ws.columns:
                max_length = 0
                column = [cell for cell in column]
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column[0].column_letter].width = min(adjusted_width, 100)
            
            # Save workbook
            wb.save(filepath)
            logger.info(f"Exported audit trail to {filepath}")
            
        except Exception as e:
            logger.error(f"Failed to export audit trail: {str(e)}")
            raise
    
    def log_allocation(self, allocation: ShiftAllocation) -> None:
        """Log a shift allocation with all compliance metrics"""
        date_str = allocation.shift.date.strftime("%Y-%m-%d")
        log_file = self.log_dir / f"{date_str}.json"
        
        # Load existing logs for the day
        allocations = []
        if log_file.exists():
            with open(log_file, 'r') as f:
                allocations = json.load(f)
        
        # Add new allocation
        allocations.append({
            "timestamp": allocation.timestamp.isoformat(),
            "doctor_id": allocation.doctor.id,
            "doctor_name": allocation.doctor.name,
            "shift_id": allocation.shift.id,
            "shift_type": allocation.shift.shift_type,
            "gender_compliant": allocation.gender_compliant,
            "team_compliant": allocation.team_compliant,
            "total_shifts": allocation.equity_stats.total_shifts,
            "night_shifts": allocation.equity_stats.night_shifts,
            "weekend_shifts": allocation.equity_stats.weekend_shifts
        })
        
        # Save updated logs
        with open(log_file, 'w') as f:
            json.dump(allocations, f, indent=2)
    
    def generate_daily_report(self, date: date) -> bytes:
        """Generate Excel report for all allocations on a specific date"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Daily Schedule"
        
        # Set up headers
        headers = ["Time", "Doctor", "Shift Type", "Gender Compliant", 
                  "Team Compliant", "Total Shifts", "Night Shifts", "Weekend Shifts"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Load and write allocation data
        log_file = self.log_dir / f"{date.strftime('%Y-%m-%d')}.json"
        row = 2
        if log_file.exists():
            with open(log_file, 'r') as f:
                allocations = json.load(f)
                for alloc in allocations:
                    ws.cell(row=row, column=1).value = datetime.fromisoformat(alloc["timestamp"]).strftime("%H:%M")
                    ws.cell(row=row, column=2).value = alloc["doctor_name"]
                    ws.cell(row=row, column=3).value = alloc["shift_type"]
                    ws.cell(row=row, column=4).value = "Yes" if alloc["gender_compliant"] else "No"
                    ws.cell(row=row, column=5).value = "Yes" if alloc["team_compliant"] else "No"
                    ws.cell(row=row, column=6).value = alloc["total_shifts"]
                    ws.cell(row=row, column=7).value = alloc["night_shifts"]
                    ws.cell(row=row, column=8).value = alloc["weekend_shifts"]
                    row += 1
        
        # Save to bytes
        from io import BytesIO
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    
    def generate_distribution_report(self, start_date: date, end_date: date) -> bytes:
        """Generate Excel report showing shift distribution metrics"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Distribution Report"
        
        # Set up headers
        headers = ["Doctor", "Total Shifts", "Night Shifts", "Weekend Shifts", 
                  "Gender Compliance %", "Team Compliance %"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Aggregate data across date range
        doctor_stats: Dict[str, Dict] = {}
        current_date = start_date
        while current_date <= end_date:
            log_file = self.log_dir / f"{current_date.strftime('%Y-%m-%d')}.json"
            if log_file.exists():
                with open(log_file, 'r') as f:
                    allocations = json.load(f)
                    for alloc in allocations:
                        doctor_id = alloc["doctor_id"]
                        if doctor_id not in doctor_stats:
                            doctor_stats[doctor_id] = {
                                "name": alloc["doctor_name"],
                                "total_shifts": 0,
                                "night_shifts": 0,
                                "weekend_shifts": 0,
                                "gender_compliant": 0,
                                "team_compliant": 0,
                                "total_allocations": 0
                            }
                        
                        stats = doctor_stats[doctor_id]
                        stats["total_shifts"] += 1
                        stats["night_shifts"] += alloc["night_shifts"]
                        stats["weekend_shifts"] += alloc["weekend_shifts"]
                        stats["gender_compliant"] += 1 if alloc["gender_compliant"] else 0
                        stats["team_compliant"] += 1 if alloc["team_compliant"] else 0
                        stats["total_allocations"] += 1
            
            current_date = date.fromordinal(current_date.toordinal() + 1)
        
        # Write aggregated data
        row = 2
        for doctor_id, stats in doctor_stats.items():
            ws.cell(row=row, column=1).value = stats["name"]
            ws.cell(row=row, column=2).value = stats["total_shifts"]
            ws.cell(row=row, column=3).value = stats["night_shifts"]
            ws.cell(row=row, column=4).value = stats["weekend_shifts"]
            ws.cell(row=row, column=5).value = f"{(stats['gender_compliant'] / stats['total_allocations']) * 100:.1f}%"
            ws.cell(row=row, column=6).value = f"{(stats['team_compliant'] / stats['total_allocations']) * 100:.1f}%"
            row += 1
        
        # Save to bytes
        from io import BytesIO
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue() 